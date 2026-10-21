import io
import json
import threading
import time
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from test_day_review import add_source
from test_extraction import sample
from test_ingest import pdf, ws  # noqa: F401

from food_record_digitizer import db
from food_record_digitizer.desktop_export import export_archive
from food_record_digitizer.desktop_import import commit_plan, import_plan
from food_record_digitizer.desktop_queue import DEFAULTS, DesktopQueue
from food_record_digitizer.desktop_server import Settings, study_app
from food_record_digitizer.imaging import render_pages
from food_record_digitizer.ingest import InputError


def wait_queue(q):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        if not q.status()["active"] and not q.status()["queued"]:
            return
        time.sleep(0.05)
    pytest.fail("queue did not finish")


def test_existing_candidates_reused_across_config_changes(ws, tmp_path):  # noqa: F811
    add_source(ws, tmp_path)
    q = DesktopQueue(ws, lambda config: pytest.fail("must not call provider"))
    try:
        p = q.preview(["R001"])
        assert p["reuse"] == 1 and p["requests"] == 0
        q.enqueue(["R001"], {**DEFAULTS, "model": "changed-model"}, p["versions"])
        assert q.status()["queued"] == 0
        with ws.factory() as s:
            assert list(s.scalars(select(db.ExtractionRun))) == []
    finally:
        q.close()


def test_double_click_dispatches_once_and_timeout_requires_confirmation(ws, tmp_path, monkeypatch):  # noqa: F811
    add_source(ws, tmp_path, candidate=False)
    calls = []

    class Provider:
        def close(self):
            pass

    def extract(*args, **kwargs):
        calls.append(1)
        time.sleep(0.15)
        return {"status": "failed", "error": "request_timeout_outcome_unknown"}

    monkeypatch.setattr("food_record_digitizer.desktop_queue.extract_record", extract)
    q = DesktopQueue(ws, lambda config: Provider())
    try:
        p = q.preview(["R001"])
        q.enqueue(["R001"], DEFAULTS, p["versions"])
        q.enqueue(["R001"], DEFAULTS, p["versions"])
        wait_queue(q)
        assert len(calls) == 1
        assert q.records()[0]["state"] == "needs_attention"
        with pytest.raises(InputError, match="charged"):
            q.enqueue(["R001"], DEFAULTS, p["versions"])
    finally:
        q.close()
    reopened = DesktopQueue(ws, lambda config: pytest.fail("must not resume automatically"))
    try:
        assert reopened.paused
        assert reopened.records()[0]["state"] == "needs_attention"
    finally:
        reopened.close()


def test_interrupted_queue_reopens_paused(ws, tmp_path):  # noqa: F811
    add_source(ws, tmp_path, candidate=False)
    q = DesktopQueue(ws)
    with q.connection() as c:
        c.execute(
            "INSERT INTO jobs (record_id,version,state,config) VALUES ('R001',1,'processing',?)",
            (json.dumps(DEFAULTS),),
        )
    q.close()
    q = DesktopQueue(ws)
    try:
        assert q.records()[0]["state"] == "needs_attention"
        assert q.paused
    finally:
        q.close()


def test_live_claim_cannot_be_recovered_or_dispatched_twice(ws, tmp_path, monkeypatch):  # noqa: F811
    add_source(ws, tmp_path, candidate=False)
    claimed, finish = threading.Event(), threading.Event()
    calls = []

    class Provider:
        def close(self):
            raise RuntimeError("transport cleanup failed")

    def extract(*args, **kwargs):
        calls.append(1)
        with ws.factory() as s:
            record = s.scalar(select(db.Record))
            record.active_run_id = "live-request"
            s.commit()
        claimed.set()
        assert finish.wait(5)
        with ws.factory() as s:
            record = s.scalar(select(db.Record))
            record.active_run_id = None
            s.commit()
        return {"status": "failed", "error": "request_timeout_outcome_unknown"}

    monkeypatch.setattr("food_record_digitizer.desktop_queue.extract_record", extract)
    monkeypatch.setattr(
        "food_record_digitizer.desktop_queue.recover_run",
        lambda *args: pytest.fail("must not recover a live request"),
    )
    q = DesktopQueue(ws, lambda config: Provider())
    try:
        preview = q.preview(["R001"])
        q.enqueue(["R001"], DEFAULTS, preview["versions"])
        assert claimed.wait(5)
        live = q.preview(["R001"])
        assert live["records"][0]["state"] == "processing"
        assert live["requests"] == 0 and live["unknown"] == []
        q.enqueue(["R001"], DEFAULTS, preview["versions"], retry_unknown=True)
        finish.set()
        wait_queue(q)
        assert calls == [1]
        assert q.records()[0]["state"] == "needs_attention"
    finally:
        finish.set()
        q.close()


def test_editable_pairing_no_rename_duplicate_rejected(ws, tmp_path):  # noqa: F811
    folder = tmp_path / "upload"
    folder.mkdir()
    main = pdf(folder / "file-0.pdf")
    attachment = pdf(folder / "file-1.pdf")
    (folder / "upload.json").write_text(json.dumps({"files": [main.name, attachment.name]}))
    entries = [{"file": main.name, "source_id": "R100"}, {"file": attachment.name, "source_id": "R100-A1"}]
    rows = [{"record_id": "R100", "subject_id": "P1", "date": None}]
    first = commit_plan(ws, folder, entries, rows)
    second = commit_plan(ws, folder, entries, rows)
    assert first["changed"] == ["R100"] and second["unchanged"] == ["R100"]
    assert main.exists() and attachment.exists()
    entries[1]["source_id"] = "R100"
    with pytest.raises(InputError, match="Duplicate"):
        import_plan(folder, entries, rows)


def test_capability_origin_no_password_and_exports(ws, tmp_path):  # noqa: F811
    add_source(ws, tmp_path)
    app = study_app(ws, "test-capability")
    try:
        client = TestClient(app)
        assert client.get("/desktop/state").status_code == 403
        client.headers["x-desktop-capability"] = "test-capability"
        assert client.get("/desktop/state", headers={"Origin": "https://bad.example"}).status_code == 403
        assert client.get("/desktop/state").json()["records"][0]["has_candidate"]
        me = client.get("/api/me").json()
        assert me["desktop"] and me["role"] == "owner"
        assert client.get("/review").status_code == 200
        assert client.post("/desktop/process/start", json={"ids": ["R001"]}).status_code == 400
        archive = client.post("/desktop/export/drafts", json={"ids": ["R001"]})
        assert archive.status_code == 200
        with zipfile.ZipFile(io.BytesIO(archive.content)) as z:
            assert not json.loads(z.read("manifest.json"))["research_approved"]
            records = json.loads(z.read("records.json"))
            workbook = load_workbook(io.BytesIO(z.read("records.xlsx")))
            assert workbook["Food items"].max_row - 1 == len(records[0]["content"]["items"])
    finally:
        app.state.queue.close()


def test_excel_evidence_not_cartesian_and_formula_safe():
    from food_record_digitizer.schemas import RecordContent

    content = RecordContent.model_validate(sample()).model_dump(mode="json")
    content["items"][0]["description"]["normalized"] = "=bad()"
    content["items"][0]["evidence"] *= 2
    data = export_archive([{"record_id": "R001", "status": "ai_draft", "content": content}])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        book = load_workbook(io.BytesIO(archive.read("records.xlsx")))
        assert book["Food items"].max_row == 2
        assert any(c.value == "'=bad()" for row in book["Food items"] for c in row)
        assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
        assert book["Evidence"].max_row > 2


def test_bundled_renderer_without_poppler(ws, tmp_path, monkeypatch):  # noqa: F811
    monkeypatch.setenv("FOOD_RECORD_RENDERER", "pdfium")
    monkeypatch.setenv("PATH", "/no-tools")
    document = {**ws.store.put(pdf(tmp_path / "R1.pdf")), "role": "main"}
    rows = render_pages(ws.store, {"documents": {"R1": document}}, tmp_path / "rendered")
    assert rows[0]["path"].read_bytes().startswith(b"\x89PNG")


def test_settings_reject_external_key_destination():
    with pytest.raises(InputError):
        Settings(base_url="https://untrusted.example/v1").checked()


def test_reopen_preserves_release_and_previous_submission(ws, tmp_path):  # noqa: F811
    from test_day_review import finish

    from food_record_digitizer.review import DEFAULT_POLICY

    add_source(ws, tmp_path)
    app = study_app(ws, "secret")
    try:
        client = TestClient(app, headers={"x-desktop-capability": "secret"})
        owner = client.get("/api/me").json()["id"]
        service = app.state.review
        service.set_policy(owner, {**DEFAULT_POLICY, "sample_percent": 0, "high_risk_flags": False})
        rid = service.queue(owner)[0]["record_id"]
        task = service.assign(owner, rid, owner)
        finish(service, owner, task)
        service.approve(owner, rid)
        release_id = service.release(owner, [rid])
        with ws.factory() as s:
            original = json.dumps(s.get(db.ReviewRelease, release_id).payload, sort_keys=True)
            revision = s.get(db.DayReviewTask, task).revision
        result = client.post("/api/tasks/" + task + "/reopen", json={"revision": revision})
        assert result.status_code == 200, result.text
        with ws.factory() as s:
            assert json.dumps(s.get(db.ReviewRelease, release_id).payload, sort_keys=True) == original
            assert s.get(db.DayReviewTask, task).status == "returned"
            assert s.get(db.DayReviewTask, result.json()["task_id"]).status == "assigned"
            unit = s.scalar(select(db.DayReviewUnit).where(db.DayReviewUnit.record_id == rid))
            assert unit.approved_id is None
        assert client.post("/api/tasks/" + task + "/reopen", json={"revision": revision}).status_code == 409
    finally:
        app.state.queue.close()


def test_manual_draft_no_ai_and_backup_restore(ws, tmp_path):  # noqa: F811
    from food_record_digitizer.backup import create_backup, restore_backup
    from food_record_digitizer.workspace import Workspace

    add_source(ws, tmp_path, candidate=False)
    app = study_app(ws, "secret")
    try:
        client = TestClient(app, headers={"x-desktop-capability": "secret"})
        assert client.post("/desktop/manual/R001", json={}).status_code == 200
        assert client.post("/desktop/manual/R001", json={}).json()["reused"]
        assert client.get("/api/records").json()[0]["has_candidate"]
        with ws.factory() as s:
            assert not list(s.scalars(select(db.ExtractionRun)))
    finally:
        app.state.queue.close()
    backup = tmp_path / "backup.zip"
    create_backup(ws, backup)
    restored = tmp_path / "restored"
    restore_backup(backup, restored)
    assert Workspace(restored).records()[0].record_id == "R001"
    assert (restored / "desktop-jobs.sqlite").exists()


@pytest.mark.parametrize(
    "error", ["api_http_401", "api_http_429", "schema_validation_failed", "output_budget_exhausted"]
)
def test_confirmed_failure_retained_and_can_retry(ws, tmp_path, monkeypatch, error):  # noqa: F811
    add_source(ws, tmp_path, candidate=False)

    class Provider:
        def close(self):
            pass

    monkeypatch.setattr(
        "food_record_digitizer.desktop_queue.extract_record",
        lambda *a, **k: {"status": "failed", "error": error},
    )
    q = DesktopQueue(ws, lambda config: Provider())
    try:
        p = q.preview(["R001"])
        q.enqueue(["R001"], DEFAULTS, p["versions"])
        wait_queue(q)
        assert q.records()[0]["state"] == "failed"
        assert q.records()[0]["error"] == error
        assert not q.preview(["R001"])["unknown"]
    finally:
        q.close()


def test_queue_persists_run_before_provider_request(ws, tmp_path, monkeypatch):  # noqa: F811
    from test_extraction import FakeProvider

    add_source(ws, tmp_path, candidate=False)
    monkeypatch.setattr("food_record_digitizer.extraction.render_pages", lambda *args: [])
    checked = []

    class Provider(FakeProvider):
        def extract(self, record_id, pages):
            with q.connection() as c:
                run_id, state = c.execute(
                    "SELECT run_id,state FROM jobs WHERE record_id=?", (record_id,)
                ).fetchone()
            with ws.factory() as s:
                record = s.scalar(select(db.Record))
                assert run_id == record.active_run_id
                assert s.get(db.ExtractionRun, run_id).status == "running"
            assert state == "processing"
            checked.append(run_id)
            return super().extract(record_id, pages)

        def close(self):
            pass

    q = DesktopQueue(ws, lambda config: Provider())
    try:
        p = q.preview(["R001"])
        q.enqueue(["R001"], DEFAULTS, p["versions"])
        wait_queue(q)
        assert len(checked) == 1
        assert q.records()[0]["run_id"] == checked[0]
        assert q.records()[0]["has_candidate"]
    finally:
        q.close()
    reopened = DesktopQueue(ws)
    try:
        assert reopened.records()[0]["run_id"] == checked[0]
        assert reopened.paused
    finally:
        reopened.close()


def test_killed_worker_reopens_unknown_without_resending(ws, tmp_path):  # noqa: F811
    import subprocess
    import sys
    from pathlib import Path

    add_source(ws, tmp_path, candidate=False)
    marker = tmp_path / "request-started"
    script = """
import sys, time
from pathlib import Path
from food_record_digitizer.workspace import Workspace
from food_record_digitizer.desktop_queue import DesktopQueue, DEFAULTS
import food_record_digitizer.extraction as extraction
extraction.render_pages = lambda *args: []
class Provider:
    model = "synthetic"
    base_url = "https://example.invalid"
    max_output_tokens = 1000
    def extract(self, record_id, pages):
        Path(sys.argv[2]).write_text("request started")
        time.sleep(120)
    def close(self): pass
q = DesktopQueue(Workspace(Path(sys.argv[1])), lambda config: Provider())
p = q.preview(["R001"])
q.enqueue(["R001"], DEFAULTS, p["versions"])
time.sleep(120)
"""
    process = subprocess.Popen(
        [sys.executable, "-c", script, str(ws.path), str(marker)], cwd=Path(__file__).parents[1]
    )
    try:
        deadline = time.monotonic() + 10
        while not marker.exists() and time.monotonic() < deadline:
            assert process.poll() is None
            time.sleep(0.05)
        assert marker.exists()
        process.kill()
        process.wait(timeout=5)
    finally:
        if process.poll() is None:
            process.kill()
            process.wait(timeout=5)
    q = DesktopQueue(ws, lambda config: pytest.fail("must not send after restart"))
    try:
        row = q.records()[0]
        assert row["state"] == "needs_attention"
        assert row["run_id"]
        assert q.status() == {"paused": True, "active": 0, "queued": 0}
        with ws.factory() as s:
            assert s.scalar(select(db.Record)).active_run_id == row["run_id"]
        preview = q.preview(["R001"])
        with pytest.raises(InputError, match="charged"):
            q.enqueue(["R001"], DEFAULTS, preview["versions"])
    finally:
        q.close()


def test_import_preview_lists_affected_days_without_changing_study(ws, tmp_path):  # noqa: F811
    from food_record_digitizer.day_review import DayReviewService
    from food_record_digitizer.desktop_import import preview_plan
    from food_record_digitizer.desktop_server import identity

    add_source(ws, tmp_path)
    add_source(ws, tmp_path, rid="R002", sid="NEW", date="2026-10-20")
    service = DayReviewService(ws)
    owner = identity(ws)["id"]
    service.queue(owner)
    with ws.factory() as s:
        before = [(r.id, r.input_version, r.fingerprint) for r in s.scalars(select(db.Record))]
    directory = tmp_path / "preview-input"
    directory.mkdir()
    (directory / "upload.json").write_text('{"files": []}')
    # Metadata alone is enough to regroup an existing source.
    result = preview_plan(
        ws, directory, [], [{"record_id": "R001", "subject_id": "NEW", "date": "2026-10-20"}]
    )
    assert result["changed"] == ["R001"]
    assert len(result["affected_reviews"]) == 2
    assert {tuple(u["records"]) for u in result["affected_reviews"]} == {("R001",), ("R002",)}
    with ws.factory() as s:
        assert [(r.id, r.input_version, r.fingerprint) for r in s.scalars(select(db.Record))] == before
        record = s.scalar(select(db.Record).where(db.Record.record_id == "R001"))
        from food_record_digitizer.ingest import current_snapshot

        original = current_snapshot(s, record)["metadata"]["R001"]
    same = preview_plan(ws, directory, [], [original])
    assert same["changed"] == [] and same["unchanged"] == ["R001"]
    assert same["affected_reviews"] == []
    assert list(directory.glob("preview-*")) == []


@pytest.mark.parametrize("error", ["api_http_401", "api_http_403", "api_http_429"])
@pytest.mark.parametrize("recovery", ["resume", "retry"])
def test_shared_api_failure_pauses_unsent_jobs(ws, tmp_path, monkeypatch, error, recovery):  # noqa: F811
    add_source(ws, tmp_path, candidate=False)
    add_source(ws, tmp_path, rid="R002", candidate=False)
    calls, configs = [], []

    class Provider:
        def close(self):
            pass

    def fail(*args, **kwargs):
        calls.append(args[1])
        return {"status": "failed", "error": error} if len(calls) == 1 else {"status": "cached"}

    monkeypatch.setattr("food_record_digitizer.desktop_queue.extract_record", fail)
    q = DesktopQueue(ws, lambda config: configs.append(dict(config)) or Provider())
    try:
        preview = q.preview(["R001", "R002"])
        q.enqueue(["R001", "R002"], {**DEFAULTS, "concurrency": 1}, preview["versions"])
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status = q.status()
            if status["paused"] and not status["active"]:
                break
            time.sleep(0.02)
        assert q.status() == {"paused": True, "active": 0, "queued": 1}
        assert len(calls) == 1
        if recovery == "resume":
            q.resume("https://us.api.openai.com/v1")
        else:
            retry = q.preview([calls[0]])
            q.enqueue(
                [calls[0]],
                {**DEFAULTS, "concurrency": 1, "base_url": "https://us.api.openai.com/v1"},
                retry["versions"],
            )
        wait_queue(q)
        assert len(calls) == (2 if recovery == "resume" else 3)
        assert all(c["base_url"] == "https://us.api.openai.com/v1" for c in configs[1:])
        assert configs[0]["base_url"] == DEFAULTS["base_url"]
        assert configs[1]["base_url"] == "https://us.api.openai.com/v1"
        assert configs[0]["model"] == configs[1]["model"]
    finally:
        q.close()


def test_desktop_records_exposes_safe_regional_diagnostic(ws, tmp_path):  # noqa: F811
    add_source(ws, tmp_path, candidate=False)
    diagnostic = {"provider_code": "incorrect_hostname", "required_base_url": "https://us.api.openai.com/v1"}
    with ws.factory.begin() as session:
        record = session.scalar(select(db.Record))
        session.add(
            db.ExtractionRun(
                record_pk=record.id,
                input_version=record.input_version,
                cache_key="synthetic-region-check",
                status="failed",
                error_code="api_http_401",
                configuration={"failure_diagnostic": diagnostic},
            )
        )
    q = DesktopQueue(ws)
    try:
        row = q.records()[0]
        assert row["error"] == "api_http_401"
        assert row["failure_diagnostic"] == diagnostic
        assert q.status()["active"] == 0
    finally:
        q.close()
