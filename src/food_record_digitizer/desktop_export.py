"""Flat research tables with lossless JSON alongside explicit spreadsheet provenance."""

import csv
import io
import json
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font


def tables(records):
    result = {
        name: [] for name in ("Records", "Food items", "Events", "Days", "Sources", "Evidence", "Issues")
    }
    for r in records:
        rid = r["record_id"]
        content = r["content"]
        base = {
            "Record ID": rid,
            "Status": r.get("status", "approved"),
            "Participant": r.get("subject_id"),
            "Metadata date": r.get("date"),
        }
        result["Records"].append(
            {
                **base,
                "Candidate ID": r.get("candidate_id"),
                "Approval ID": r.get("approval_id"),
                "Input version": r.get("input_version"),
            }
        )
        events = {e["event_id"]: e for e in content["events"]}

        def flatten(obj, entity, kind):
            row = {**base, "Entity ID": entity}
            for key, value in obj.items():
                if key == "evidence":
                    evidence(value, kind, entity, "item")
                elif isinstance(value, dict) and "normalized" in value:
                    label = key.replace("_", " ").title()
                    row[label] = value.get("normalized")
                    row[label + " raw"] = value.get("raw")
                    row[label + " missing reason"] = value.get("missing_reason")
                    row[label + " inherited from"] = value.get("inherited_from")
                    meta = value.get("metadata_reference") or {}
                    row[label + " metadata record"] = meta.get("record_id")
                    row[label + " metadata field"] = meta.get("field")
                    evidence(value.get("evidence", []), kind, entity, key)
                elif not isinstance(value, (dict, list)):
                    row[key.replace("_", " ").title()] = value
            return row

        def evidence(values, kind, entity, field):
            for i, e in enumerate(values):
                result["Evidence"].append(
                    {
                        "Record ID": rid,
                        "Entity type": kind,
                        "Entity ID": entity,
                        "Field": field,
                        "Evidence index": i + 1,
                        "Source ID": e["source_record_id"],
                        "Page": e["page_number"],
                        "Region": e["region"],
                        "Transcription": e.get("transcription_raw"),
                        **dict(zip(("Left", "Top", "Right", "Bottom"), e.get("bbox") or [None] * 4)),
                    }
                )

        for item in content["items"]:
            row = flatten(item, item["item_id"], "food")
            event = events.get(item["event_id"], {})
            for key in ("time", "meal_label"):
                value = event.get(key, {})
                row[key.replace("_", " ").title()] = value.get("normalized")
                row[key.replace("_", " ").title() + " missing reason"] = value.get("missing_reason")
            result["Food items"].append(row)
        for kind, key, ident in (("Events", "events", "event_id"), ("Days", "days", "day_id")):
            for obj in content[key]:
                result[kind].append(flatten(obj, obj[ident], kind))
        for i, issue in enumerate(content["issues"]):
            result["Issues"].append(flatten(issue, str(i + 1), "issue"))
        for sid, doc in r.get("snapshot", r.get("sources", {})).get("documents", {}).items():
            result["Sources"].append(
                {
                    "Record ID": rid,
                    "Source ID": sid,
                    "Role": doc["role"],
                    "Original filename": doc.get("original_filename"),
                    "Study source file": doc.get("managed_file"),
                    "Attachment origin": doc.get("attachment_original", {}).get("origin"),
                    "Attachment added at": doc.get("attachment_original", {}).get("added_at"),
                    "Original attachment SHA256": doc.get("attachment_original", {}).get("sha256"),
                    "Pages": doc["page_count"],
                    "SHA256": doc["sha256"],
                }
            )
    return result


def cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def column_meaning(column):
    meanings = {
        "Record ID": "Stable source record identifier; use with entity IDs when joining tables.",
        "Status": "ai_draft, manual_draft or approved; drafts are not approved research data.",
        "Participant": "Participant identifier from authoritative metadata.",
        "Metadata date": "Intake date from metadata; blank means unavailable, not an inferred date.",
        "Entity ID": "Identifier within this record and entity type.",
        "Event Id": "Join to Events using Record ID and Event Id.",
        "Day Id": "Join to Days using Record ID and Day Id.",
        "Candidate ID": "Identifier of the extracted or manually created draft.",
        "Approval ID": "Identifier of the approved version, when applicable.",
        "Input version": "Revision of the source documents and metadata.",
        "Entity type": "Evidence owner: food, Events, Days or issue.",
        "Field": "Field on the evidence owner; item denotes whole-item evidence.",
        "Evidence index": "One-based reference index within the owner field.",
        "Source ID": "Main document or attachment identifier; join to Sources by Record ID and Source ID.",
        "Page": "One-based page number within the source document.",
        "Pages": "Total pages in the source document.",
        "Region": "Human-readable location on the scanned page.",
        "Transcription": "Original text supporting the field.",
        "Role": "Source document role: main or attachment.",
        "Original filename": "Original input filename before the app assigned a source ID.",
        "Study source file": "Managed copy path relative to the owner study sources folder. Original source bytes are also identified by SHA256.",
        "Attachment origin": "participant_material or researcher_note; distinguishes original participant material from later researcher explanation.",
        "Attachment added at": "UTC time the attachment was staged by the study owner.",
        "Original attachment SHA256": "SHA-256 of the original attachment bytes before PDF conversion.",
        "SHA256": "Source file SHA-256 digest for content verification.",
        "Time": "Recorded time, independent of meal label; blank remains unknown.",
        "Meal Label": "Recorded meal label, independent of time; blank remains unknown.",
    }
    if column in meanings:
        return meanings[column]
    if column in {"Left", "Top", "Right", "Bottom"}:
        return "Bounding-box coordinate normalized to 0–1 on the page; blank if unavailable."
    for suffix, explanation in {
        " missing reason": "Reason the normalized value is absent; absence is never zero.",
        " inherited from": "Identifier of the source entity from which the value was inherited.",
        " metadata record": "Record ID of the metadata supporting the value.",
        " metadata field": "Metadata field supporting the value.",
        " raw": "Original recorded text before normalization.",
    }.items():
        if column.endswith(suffix):
            return f"{column.removesuffix(suffix)}: {explanation}"
    return f"{column}: structured field from records.json; blank indicates no value."


def export_archive(records, approved=False, release_id=None):
    data = tables(records)
    book = Workbook()
    book.remove(book.active)
    dictionary = book.create_sheet("Read me")
    dictionary.append(["Property", "Meaning"])
    for row in [
        ("Status", "Approved release" if approved else "Draft — not research approved"),
        ("Release ID", release_id),
        ("Blank cells", "Unknown or not recorded; never interpreted as zero"),
        ("Evidence", "Join Record ID + Entity type + Entity ID + Field; one row per evidence reference"),
        ("Food time / label", "Separate values linked by Event ID. Neither implies the other."),
        ("JSON", "records.json preserves complete content, provenance and nested references"),
        (
            "Formula protection",
            "Text starting with =, +, -, @, tab or carriage return has a leading apostrophe in CSV/XLSX; JSON retains the original text.",
        ),
    ]:
        dictionary.append(row)
    fields = book.create_sheet("Data dictionary")
    fields.append(["Table", "Column", "Meaning"])
    csvs = {}
    for name, rows in data.items():
        columns = list(dict.fromkeys(k for row in rows for k in row)) or ["Record ID"]
        for column in columns:
            fields.append([name, column, column_meaning(column)])
        sheet = book.create_sheet(name)
        sheet.append(columns)
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(columns)
        for row in rows:
            values = [cell(row.get(k)) for k in columns]
            sheet.append(values)
            writer.writerow(values)
        csvs[name.lower().replace(" ", "_") + ".csv"] = stream.getvalue()
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for c in sheet[1]:
            c.font = Font(bold=True)
            sheet.column_dimensions[c.column_letter].width = min(45, max(18, len(str(c.value)) + 2))
    excel = io.BytesIO()
    book.save(excel)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("records.xlsx", excel.getvalue())
        archive.writestr("records.json", json.dumps(records, ensure_ascii=False, indent=2))
        archive.writestr(
            "manifest.json",
            json.dumps(
                {"research_approved": approved, "release_id": release_id, "record_count": len(records)}
            ),
        )
        for name, value in csvs.items():
            archive.writestr(name, value)
    return output.getvalue()
